"""
Document Parser and Text Chunker for DigitalBrainEX AI.
Extracts structured text from PDFs (page-by-page), Word documents (docx sections),
and plaintext files, then splits them into overlapping chunks with location metadata.
"""
import os
import sys
import re
import csv
from typing import List, Dict, Any, Tuple
from src.core.logger import logger


class DocumentParser:
    """Extracts text and produces overlapping chunks with page/section citations."""

    DEFAULT_CHUNK_SIZE = 1000     # characters per chunk (~150-200 words)
    DEFAULT_CHUNK_OVERLAP = 200   # overlap between adjacent chunks

    @classmethod
    def parse_and_chunk(
        cls,
        file_path: str,
        chunk_size: int = DEFAULT_CHUNK_SIZE,
        chunk_overlap: int = DEFAULT_CHUNK_OVERLAP,
    ) -> List[Dict[str, Any]]:
        """
        Parses the document at file_path and splits it into indexed chunks.
        Supports PDF, DOCX, DOC, PPTX, PPT, RTF, XLSX, CSV, and plaintext formats.
        Returns a list of dicts: [{'chunk_index': int, 'chunk_text': str, 'page_or_section': str}]
        """
        if not file_path or not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".pdf":
            sections = cls._extract_pdf(file_path)
        elif ext == ".docx":
            sections = cls._extract_docx(file_path)
        elif ext == ".doc":
            sections = cls._extract_doc(file_path)
        elif ext in (".pptx", ".pptm"):
            sections = cls._extract_pptx(file_path)
        elif ext == ".ppt":
            sections = cls._extract_ppt(file_path)
        elif ext == ".rtf":
            sections = cls._extract_rtf(file_path)
        elif ext in (".xlsx", ".xlsm", ".xls"):
            sections = cls._extract_xlsx(file_path)
        elif ext in (".csv", ".tsv"):
            sections = cls._extract_csv(file_path)
        else:
            sections = cls._extract_plaintext(file_path)

        if not sections:
            return []

        # Chunk the extracted sections
        all_chunks = []
        global_chunk_idx = 0

        for page_or_sec, text in sections:
            clean_text = text.strip()
            if not clean_text:
                continue

            section_chunks = cls._split_into_chunks(clean_text, chunk_size, chunk_overlap)
            for ch in section_chunks:
                all_chunks.append({
                    "chunk_index": global_chunk_idx,
                    "chunk_text": ch,
                    "page_or_section": page_or_sec,
                })
                global_chunk_idx += 1

        return all_chunks

    @classmethod
    def _extract_pdf(cls, file_path: str) -> List[Tuple[str, str]]:
        """Extracts text per page using PyPDF2."""
        sections = []
        try:
            import PyPDF2
            with open(file_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                num_pages = len(reader.pages)
                for page_idx in range(num_pages):
                    page = reader.pages[page_idx]
                    page_text = page.extract_text() or ""
                    if page_text.strip():
                        sections.append((f"Page {page_idx + 1}", page_text))
        except Exception as e:
            logger.error(f"Error parsing PDF '{file_path}': {e}")
            raise RuntimeError(f"Could not extract text from PDF: {e}") from e

        return sections

    @classmethod
    def _extract_docx(cls, file_path: str) -> List[Tuple[str, str]]:
        """Extracts text per section/heading using python-docx."""
        sections = []
        try:
            import docx
            doc = docx.Document(file_path)
            current_heading = "Introduction"
            current_paragraphs = []

            for p in doc.paragraphs:
                p_text = p.text.strip()
                if not p_text:
                    continue

                # Check if paragraph is a heading
                if p.style and p.style.name and p.style.name.startswith("Heading"):
                    if current_paragraphs:
                        sections.append((f"Section: {current_heading}", "\n\n".join(current_paragraphs)))
                        current_paragraphs = []
                    current_heading = p_text[:60]
                else:
                    current_paragraphs.append(p_text)

            if current_paragraphs:
                sections.append((f"Section: {current_heading}", "\n\n".join(current_paragraphs)))

            # Also check tables in docx
            for t_idx, table in enumerate(doc.tables, 1):
                table_rows = []
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        table_rows.append(" | ".join(cells))
                if table_rows:
                    sections.append((f"Table {t_idx}", "\n".join(table_rows)))

        except Exception as e:
            logger.error(f"Error parsing DOCX '{file_path}': {e}")
            raise RuntimeError(f"Could not extract text from Word document: {e}") from e

        return sections

    @classmethod
    def _extract_plaintext(cls, file_path: str) -> List[Tuple[str, str]]:
        """Reads plain text / code / markdown files with multiple encoding fallbacks."""
        encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
        content = None
        for enc in encodings:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    content = f.read()
                    break
            except Exception:
                continue

        if content is None:
            raise ValueError(f"Could not decode text file {file_path} with supported encodings.")

        filename = os.path.basename(file_path)
        return [(f"File: {filename}", content)]

    @classmethod
    def _extract_pptx(cls, file_path: str) -> List[Tuple[str, str]]:
        """Extracts text per slide using python-pptx."""
        sections = []
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            for slide_idx, slide in enumerate(prs.slides, 1):
                slide_texts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for p in shape.text_frame.paragraphs:
                            t = p.text.strip()
                            if t:
                                slide_texts.append(t)
                    elif shape.has_table:
                        for row in shape.table.rows:
                            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                            if row_cells:
                                slide_texts.append(" | ".join(row_cells))
                if slide_texts:
                    sections.append((f"Slide {slide_idx}", "\n".join(slide_texts)))
        except Exception as e:
            logger.warning(f"Error parsing PPTX '{file_path}': {e}. Trying binary fallback...")
            fallback = cls._extract_binary_strings(file_path)
            if fallback:
                sections.append((f"File: {os.path.basename(file_path)}", fallback))
            else:
                raise RuntimeError(f"Could not extract text from PPTX: {e}") from e
        return sections

    @classmethod
    def _extract_ppt(cls, file_path: str) -> List[Tuple[str, str]]:
        """Extracts text from legacy PowerPoint 97-2003 (.ppt) using COM automation or binary fallback."""
        sections = []
        abs_path = os.path.abspath(file_path)
        if sys.platform == "win32":
            try:
                import win32com.client
                import pythoncom
                pythoncom.CoInitialize()
                try:
                    ppt_app = win32com.client.Dispatch("PowerPoint.Application")
                    presentation = ppt_app.Presentations.Open(abs_path, WithWindow=False, ReadOnly=True)
                    for slide_idx in range(1, presentation.Slides.Count + 1):
                        slide = presentation.Slides(slide_idx)
                        slide_texts = []
                        for shape in slide.Shapes:
                            if shape.HasTextFrame and shape.TextFrame.HasText:
                                t = shape.TextFrame.TextRange.Text.strip()
                                if t:
                                    slide_texts.append(t)
                        if slide_texts:
                            sections.append((f"Slide {slide_idx}", "\n".join(slide_texts)))
                    presentation.Close()
                finally:
                    pythoncom.CoUninitialize()
                if sections:
                    return sections
            except Exception as e:
                logger.debug(f"COM extraction failed for PPT '{file_path}': {e}")

        # Fallback to binary strings extraction
        fallback = cls._extract_binary_strings(file_path)
        if fallback:
            sections.append((f"File: {os.path.basename(file_path)}", fallback))
        return sections

    @classmethod
    def _extract_doc(cls, file_path: str) -> List[Tuple[str, str]]:
        """Extracts text from Word 97-2003 (.doc) files using docx fallback, COM automation, or binary extraction."""
        # 1. Try python-docx in case it's actually OOXML with .doc extension
        try:
            docx_sections = cls._extract_docx(file_path)
            if docx_sections:
                return docx_sections
        except Exception:
            pass

        # 2. Try Word COM automation on Windows
        abs_path = os.path.abspath(file_path)
        if sys.platform == "win32":
            try:
                import win32com.client
                import pythoncom
                pythoncom.CoInitialize()
                try:
                    word_app = win32com.client.Dispatch("Word.Application")
                    word_app.Visible = False
                    doc = word_app.Documents.Open(abs_path, ReadOnly=True, Visible=False)
                    text = doc.Content.Text
                    doc.Close(SaveChanges=False)
                    if text and text.strip():
                        return [(f"File: {os.path.basename(file_path)}", text.strip())]
                finally:
                    pythoncom.CoUninitialize()
            except Exception as e:
                logger.debug(f"COM extraction failed for DOC '{file_path}': {e}")

        # 3. Fallback to binary strings extraction
        fallback = cls._extract_binary_strings(file_path)
        if fallback:
            return [(f"File: {os.path.basename(file_path)}", fallback)]
        return []

    @classmethod
    def _extract_rtf(cls, file_path: str) -> List[Tuple[str, str]]:
        """Extracts clean text from RTF files by stripping binary images and formatting tags."""
        encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
        raw_rtf = None
        for enc in encodings:
            try:
                with open(file_path, "r", encoding=enc, errors="ignore") as f:
                    raw_rtf = f.read()
                    break
            except Exception:
                continue

        if not raw_rtf:
            return []

        # Strip embedded picture hex dumps and header blocks
        clean = re.sub(r'{\\\*\\pict[\s\S]*?}', '', raw_rtf)
        clean = re.sub(r'{\\pict[\s\S]*?}', '', clean)
        clean = re.sub(r'{\\colortbl[\s\S]*?}', '', clean)
        clean = re.sub(r'{\\fonttbl[\s\S]*?}', '', clean)
        clean = re.sub(r'{\\stylesheet[\s\S]*?}', '', clean)
        clean = re.sub(r'{\\info[\s\S]*?}', '', clean)
        clean = re.sub(r'\\[a-zA-Z0-9\-]+ ?', ' ', clean)
        clean = re.sub(r'[{}\\]', ' ', clean)
        clean = ' '.join(clean.split())

        if clean:
            return [(f"File: {os.path.basename(file_path)}", clean)]
        return []

    @classmethod
    def _extract_xlsx(cls, file_path: str) -> List[Tuple[str, str]]:
        """Extracts text from Excel spreadsheets per sheet using openpyxl."""
        sections = []
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                rows_text = []
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                    if row_vals:
                        rows_text.append(" | ".join(row_vals))
                if rows_text:
                    sections.append((f"Sheet: {sheet_name}", "\n".join(rows_text)))
            wb.close()
        except Exception as e:
            logger.warning(f"Error parsing XLSX '{file_path}': {e}")
            fallback = cls._extract_binary_strings(file_path)
            if fallback:
                sections.append((f"File: {os.path.basename(file_path)}", fallback))
            else:
                raise RuntimeError(f"Could not extract text from Excel workbook: {e}") from e
        return sections

    @classmethod
    def _extract_csv(cls, file_path: str) -> List[Tuple[str, str]]:
        """Extracts text from CSV and TSV tabular files."""
        encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
        for enc in encodings:
            try:
                with open(file_path, "r", encoding=enc, errors="ignore") as f:
                    delimiter = "\t" if file_path.lower().endswith(".tsv") else ","
                    reader = csv.reader(f, delimiter=delimiter)
                    lines = []
                    for row in reader:
                        cleaned = [c.strip() for c in row if c.strip()]
                        if cleaned:
                            lines.append(" | ".join(cleaned))
                    if lines:
                        return [(f"File: {os.path.basename(file_path)}", "\n".join(lines))]
            except Exception:
                continue
        return cls._extract_plaintext(file_path)

    @classmethod
    def _extract_binary_strings(cls, file_path: str, min_len: int = 4) -> str:
        """Fallback string extractor for binary files."""
        try:
            with open(file_path, "rb") as f:
                content = f.read()
            strings = re.findall(rb'[A-Za-z0-9\s.,;:?!_\-/\(\)\[\]]{%d,}' % min_len, content)
            valid = []
            for s in strings:
                decoded = s.decode('ascii', errors='ignore').strip()
                if len(decoded) >= min_len and not decoded.startswith(('!', '@', '#')):
                    valid.append(decoded)
            return "\n".join(valid[:1000])  # limit to top 1000 strings
        except Exception:
            return ""

    @classmethod
    def parse_and_chunk_text(
        cls,
        text: str,
        source_title: str = "Note",
        chunk_size: int = DEFAULT_CHUNK_SIZE,
        chunk_overlap: int = DEFAULT_CHUNK_OVERLAP,
    ) -> List[Dict[str, Any]]:
        """
        Directly chunks arbitrary text (used for PlainNotes, descriptions, or web content).
        Returns list of dicts: [{'chunk_index': int, 'chunk_text': str, 'page_or_section': str}]
        """
        if not text or not text.strip():
            return []

        clean_text = text.strip()
        section_chunks = cls._split_into_chunks(clean_text, chunk_size, chunk_overlap)
        all_chunks = []
        for idx, ch in enumerate(section_chunks):
            all_chunks.append({
                "chunk_index": idx,
                "chunk_text": ch,
                "page_or_section": source_title,
            })
        return all_chunks

    @classmethod
    def _split_into_chunks(cls, text: str, chunk_size: int, chunk_overlap: int) -> List[str]:
        """
        Splits text into chunks of at most chunk_size characters with chunk_overlap,
        respecting paragraph or sentence boundaries where possible.
        """
        if len(text) <= chunk_size:
            return [text]

        chunks = []
        start = 0
        text_len = len(text)

        while start < text_len:
            end = start + chunk_size
            if end >= text_len:
                chunks.append(text[start:].strip())
                break

            # Find natural break point near end (paragraph break or period)
            cut_point = end
            para_break = text.rfind("\n\n", start, end)
            if para_break != -1 and para_break > start + (chunk_size // 2):
                cut_point = para_break + 2
            else:
                sent_break = text.rfind(". ", start, end)
                if sent_break != -1 and sent_break > start + (chunk_size // 2):
                    cut_point = sent_break + 2
                else:
                    space_break = text.rfind(" ", start, end)
                    if space_break != -1 and space_break > start + (chunk_size // 2):
                        cut_point = space_break + 1

            chunk = text[start:cut_point].strip()
            if chunk:
                chunks.append(chunk)

            # Advance start with overlap
            start = max(start + 1, cut_point - chunk_overlap)

        return chunks
